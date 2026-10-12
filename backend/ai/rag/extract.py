"""Text extraction and validation for uploaded course documents."""

import io
import os
import re
import unicodedata
import zipfile

from config import settings

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}
_MAGIC_PDF = b"%PDF-"
_MAGIC_ZIP = b"PK\x03\x04"


class ExtractionError(ValueError):
    """The file is unsupported, corrupt, empty, or has no extractable text (user-safe message)."""


def normalize_text(text: str) -> str:
    """Tidy whitespace without changing meaning (keeps paragraph breaks, code indentation, headings)."""
    text = unicodedata.normalize("NFKC", text or "")
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)           # trailing spaces
    text = re.sub(r"\n{3,}", "\n\n", text)           # runs of blank lines
    text = re.sub(r"-\n(?=[a-z])", "", text)         # PDF hyphenation at line breaks
    return text.strip()


def validate_upload(filename: str, data: bytes) -> str:
    """Check extension, size and file signature. Returns the lower-case extension."""
    ext = os.path.splitext(filename or "")[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ExtractionError("Unsupported file type. Upload a PDF, DOCX, TXT or Markdown file.")
    if not data:
        raise ExtractionError("The file is empty.")
    if len(data) > settings.RAG_MAX_UPLOAD_MB * 1024 * 1024:
        raise ExtractionError(f"The file is larger than {settings.RAG_MAX_UPLOAD_MB} MB.")
    if ext == ".pdf" and not data.lstrip()[:5].startswith(_MAGIC_PDF):
        raise ExtractionError("The file is not a valid PDF.")
    if ext == ".docx" and not data.startswith(_MAGIC_ZIP):
        raise ExtractionError("The file is not a valid DOCX document.")
    return ext


def _pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ExtractionError("PDF support is not installed on this server.") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("Password-protected PDFs are not supported.")
        pages = []
        for number, page in enumerate(reader.pages, start=1):
            pages.append(f"[Page {number}]\n{page.extract_text() or ''}")
        return "\n\n".join(pages)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("The PDF could not be read. It may be corrupt.") from exc


def _docx_text(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:
        raise ExtractionError("DOCX support is not installed on this server.") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ExtractionError("The DOCX file could not be read. It may be corrupt.") from exc
    parts = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "") if paragraph.style is not None else ""
        if style.lower().startswith("heading"):
            parts.append(f"\n# {text}\n")
        else:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _plain_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeError:
            continue
    return data.decode("latin-1", errors="replace")


def extract_text(filename: str, data: bytes) -> str:
    """Return normalised text, or raise ExtractionError. Never returns empty text."""
    ext = validate_upload(filename, data)
    if ext == ".pdf":
        raw = _pdf_text(data)
    elif ext == ".docx":
        raw = _docx_text(data)
    else:
        raw = _plain_text(data)
    text = normalize_text(raw)
    # A scanned PDF yields only the "[Page n]" markers.
    meaningful = re.sub(r"\[Page \d+\]", "", text).strip()
    if len(meaningful) < 20:
        raise ExtractionError("No readable text was found in the file (scanned documents are not supported).")
    if len(text) > settings.RAG_MAX_DOCUMENT_CHARS:
        raise ExtractionError("The document contains too much text to index.")
    return text
