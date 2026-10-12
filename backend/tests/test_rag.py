"""
RAG pipeline tests.

Real: chunking, text extraction (real PDF/DOCX files), SQLAlchemy storage, FAISS indexes, the
authorization layer and the HTTP endpoints. Test double: the embedding MODEL is replaced by a
deterministic hashing embedder (HashEmbedder) so tests run without downloading Sentence
Transformers weights; `test_real_sentence_transformer_*` exercises the real model and is skipped
when it cannot be loaded. The LLM is replaced by a recorder so prompts can be inspected.
"""

import io
import zlib

import numpy as np
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from ai import llm
from ai.rag import chunking, embedder, extract, index, service
from ai.rag.access import RagAccessDenied, authorize_course_access
from auth.jwt_handler import create_access_token
from config import settings
from database import Base, get_db
from main import app
from models.course import Course, CourseCategory, Enrollment, Lecture, Module
from models.payment import Payment, PaymentStatus
from models.rag import CourseDocument, DocumentChunk, DocumentStatus
from models.user import User, UserRole

engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)


# --------------------------------------------------------------------------- test doubles

class HashEmbedder:
    """Bag-of-words hashing embedder: deterministic, L2-normalised, related texts share words."""
    model_name = "test/hash-embedder"
    dim = 4096

    def encode(self, texts):
        out = np.zeros((len(texts), self.dim), dtype=np.float32)
        for row, text in enumerate(texts):
            for word in "".join(c.lower() if c.isalnum() else " " for c in text).split():
                if len(word) > 2:
                    out[row, zlib.crc32(word.encode()) % self.dim] += 1.0
        norms = np.linalg.norm(out, axis=1, keepdims=True)
        norms[norms == 0] = 1.0
        return out / norms


class LLMRecorder:
    def __init__(self):
        self.calls = []
        self.reply = "Answer from the course [S1]."

    def __call__(self, prompt, **kwargs):
        self.calls.append({"prompt": prompt, **kwargs})
        return llm.LLMResult(text=self.reply, tokens_used=7)


@pytest.fixture(autouse=True)
def rag_env(monkeypatch):
    index.clear()
    embedder.set_embedder(HashEmbedder())
    monkeypatch.setattr(settings, "EMBEDDING_MODEL", HashEmbedder.model_name)
    monkeypatch.setattr(settings, "RAG_MIN_SCORE", 0.2)
    yield
    embedder.set_embedder(None)
    index.clear()


@pytest.fixture
def recorder(monkeypatch):
    rec = LLMRecorder()
    monkeypatch.setattr(llm, "generate", rec)
    return rec


@pytest.fixture
def db():
    import models  # noqa: F401
    Base.metadata.create_all(bind=engine)
    session = TestingSessionLocal()
    try:
        yield session
    finally:
        session.close()
        Base.metadata.drop_all(bind=engine)


@pytest.fixture
def client(db):
    def override():
        yield db
    app.dependency_overrides[get_db] = override
    yield TestClient(app, raise_server_exceptions=True)
    app.dependency_overrides.clear()


def auth(user):
    return {"Authorization": f"Bearer {create_access_token({'sub': str(user.id)})}"}


def make_user(db, name, role=UserRole.STUDENT):
    user = User(email=f"{name}@x.io", username=name, password_hash="x", full_name=name.title(),
                role=role, is_verified=True)
    db.add(user)
    db.flush()
    return user


PAID_TEXT = ("React hooks let function components keep state. The useState hook stores a value and the "
             "useEffect hook runs side effects after rendering. Premium secret: the capstone rubric awards "
             "marks for custom hooks.")
FREE_TEXT = "Python generators yield values lazily, which saves memory when iterating over large sequences."


@pytest.fixture
def world(db):
    instructor = make_user(db, "instructor", UserRole.INSTRUCTOR)
    other_instructor = make_user(db, "instructor2", UserRole.INSTRUCTOR)
    student = make_user(db, "student")
    stranger = make_user(db, "stranger")
    admin = make_user(db, "admin", UserRole.ADMIN)
    paid = Course(title="React Masterclass", slug="react", description=PAID_TEXT, instructor_id=instructor.id,
                  category=list(CourseCategory)[0], is_published=True, is_approved=True,
                  pricing_type="PAID", price_amount=99900, currency="INR")
    free = Course(title="Intro to Python", slug="python", description=FREE_TEXT, instructor_id=instructor.id,
                  category=list(CourseCategory)[0], is_published=True, is_approved=True)
    draft = Course(title="Unreleased Draft", slug="draft", description="Unpublished secret syllabus text here.",
                   instructor_id=instructor.id, category=list(CourseCategory)[0])
    db.add_all([paid, free, draft])
    db.flush()
    module = Module(course_id=paid.id, title="Hooks", description="Working with state", order_index=0)
    db.add(module)
    db.flush()
    db.add(Lecture(module_id=module.id, title="Introducing useState", description="Counter example", order_index=0))
    db.commit()
    return dict(instructor=instructor, other_instructor=other_instructor, student=student, stranger=stranger,
                admin=admin, paid=paid, free=free, draft=draft)


_order = iter(range(1, 10_000))


def grant_paid_access(db, user, course, status=PaymentStatus.SUCCESS.value):
    db.add(Payment(user_id=user.id, course_id=course.id, razorpay_order_id=f"order_TEST{next(_order):08d}",
                   amount=course.price_amount, currency="INR", status=status))
    db.add(Enrollment(user_id=user.id, course_id=course.id, is_active=True))
    db.commit()


def enroll(db, user, course):
    db.add(Enrollment(user_id=user.id, course_id=course.id, is_active=True))
    db.commit()


# --------------------------------------------------------------------------- chunking

def test_chunks_carry_sections_overlap_and_stable_indexes():
    body = " ".join(f"Sentence number {i} explains concept {i} in some detail." for i in range(60))
    text = f"# Intro\n{body}\n\n## Details\nShort section."
    chunks = chunking.chunk_text(text, size=300, overlap=60)
    assert [c.index for c in chunks] == list(range(len(chunks)))
    assert chunks[0].section == "Intro" and chunks[-1].section == "Details"
    assert all(len(c.text) <= 300 + 120 for c in chunks)
    intro = [c for c in chunks if c.section == "Intro"]
    assert len(intro) > 2
    # Overlap: the end of one window reappears at the start of the next.
    assert intro[0].text.split()[-2] in intro[1].text


def test_pdf_page_markers_become_sections():
    chunks = chunking.chunk_text("[Page 1]\nFirst page text here.\n\n[Page 2]\nSecond page text here.")
    assert [c.section for c in chunks] == ["Page 1", "Page 2"]


def test_single_giant_line_is_hard_split():
    chunks = chunking.chunk_text("x" * 5000, size=1000, overlap=100)
    assert len(chunks) >= 5 and all(len(c.text) <= 1000 for c in chunks)


# --------------------------------------------------------------------------- extraction

def _pdf(text):
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(72, 750, text)
    c.save()
    return buf.getvalue()


def _docx():
    import docx
    d = docx.Document()
    d.add_heading("Closures", level=1)
    d.add_paragraph("A closure captures variables from its enclosing scope.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_extract_real_pdf_docx_txt_md():
    assert "Binary search halves" in extract.extract_text("a.pdf", _pdf("Binary search halves the search space each step."))
    docx_text = extract.extract_text("b.docx", _docx())
    assert "# Closures" in docx_text and "enclosing scope" in docx_text
    assert "plain text body here" in extract.extract_text("c.txt", b"A plain text body here, long enough.")
    assert extract.extract_text("d.md", b"# Title\nMarkdown body that is long enough.").startswith("# Title")


@pytest.mark.parametrize("name,data,message", [
    ("a.exe", b"MZ....", "Unsupported"),
    ("a.txt", b"", "empty"),
    ("a.pdf", b"not a pdf at all", "not a valid PDF"),
    ("a.pdf", b"%PDF-1.4 garbage garbage garbage", "could not be read"),
    ("a.docx", b"plain bytes", "not a valid DOCX"),
    ("a.docx", b"PK\x03\x04broken", "could not be read"),
    ("a.txt", b"   \n  ", "No readable text"),
])
def test_extraction_failures_are_reported(name, data, message):
    with pytest.raises(extract.ExtractionError, match=message):
        extract.extract_text(name, data)


def test_oversized_upload_rejected(monkeypatch):
    monkeypatch.setattr(settings, "RAG_MAX_UPLOAD_MB", 1)
    with pytest.raises(extract.ExtractionError, match="larger than"):
        extract.extract_text("big.txt", b"a" * (1024 * 1024 + 1))


def test_scanned_pdf_without_text_is_rejected():
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.showPage()
    c.save()
    with pytest.raises(extract.ExtractionError, match="No readable text"):
        extract.extract_text("scan.pdf", buf.getvalue())


# --------------------------------------------------------------------------- ingestion and storage

def _ingest(db, course, title, text, owner=None):
    doc, created = service.create_document(db, course, title, text, source_type="upload",
                                           owner_id=owner.id if owner else None)
    return service.process_document(db, doc.id), created


def test_ingestion_stores_chunks_metadata_and_embedding_dimension(db, world):
    doc, created = _ingest(db, world["paid"], "Notes", PAID_TEXT * 3, world["instructor"])
    assert created and doc.status == DocumentStatus.READY.value and doc.error is None
    assert doc.embedding_model == "test/hash-embedder" and doc.embedding_dim == 4096
    rows = db.query(DocumentChunk).filter(DocumentChunk.document_id == doc.id).all()
    assert len(rows) == doc.chunk_count > 0
    assert all(r.course_id == world["paid"].id and len(r.embedding) == 4096 * 4 for r in rows)
    vec = np.frombuffer(rows[0].embedding, dtype=np.float32)
    assert abs(float(np.linalg.norm(vec)) - 1.0) < 1e-4  # normalised -> inner product == cosine
    assert doc.owner_id == world["instructor"].id and doc.visibility == "enrolled"


def test_duplicate_content_is_not_ingested_twice(db, world):
    first, created1 = _ingest(db, world["paid"], "A", PAID_TEXT)
    second, created2 = _ingest(db, world["paid"], "B", PAID_TEXT)
    assert created1 and not created2 and first.id == second.id
    assert db.query(CourseDocument).filter(CourseDocument.course_id == world["paid"].id).count() == 1
    # The same text in a different course is a different document.
    _, created3 = _ingest(db, world["free"], "A", PAID_TEXT)
    assert created3


def test_embedding_failure_marks_document_failed_not_ready(db, world, monkeypatch):
    doc, _ = service.create_document(db, world["paid"], "Doc", PAID_TEXT, source_type="upload", owner_id=None)

    def boom(texts):
        raise embedder.EmbeddingUnavailable("model could not be loaded")
    monkeypatch.setattr(embedder.get_embedder(), "encode", boom)
    with pytest.raises(embedder.EmbeddingUnavailable):
        service.process_document(db, doc.id)
    db.refresh(doc)
    assert doc.status == DocumentStatus.FAILED.value and "model could not be loaded" in doc.error
    assert db.query(DocumentChunk).filter(DocumentChunk.document_id == doc.id).count() == 0


def test_failed_reingest_keeps_previous_chunks(db, world, monkeypatch):
    doc, _ = _ingest(db, world["paid"], "Doc", PAID_TEXT)
    before = db.query(DocumentChunk).filter(DocumentChunk.document_id == doc.id).count()
    service.reset_for_reingest(db, doc)
    monkeypatch.setattr(embedder.get_embedder(), "encode", lambda texts: (_ for _ in ()).throw(RuntimeError("x")))
    with pytest.raises(RuntimeError):
        service.process_document(db, doc.id)
    assert db.query(DocumentChunk).filter(DocumentChunk.document_id == doc.id).count() == before


def test_reingest_replaces_chunks_and_delete_removes_them_from_search(db, world):
    doc, _ = _ingest(db, world["paid"], "Doc", PAID_TEXT)
    emb = embedder.get_embedder()
    query = emb.encode(["useState hook stores a value"])[0]
    assert index.search(db, world["paid"].id, query, emb.model_name, emb.dim, 3)

    doc.text = "Completely different topic: databases use B-tree indexes for fast lookups."
    service.reset_for_reingest(db, doc)
    service.process_document(db, doc.id)
    hits = index.search(db, world["paid"].id, query, emb.model_name, emb.dim, 3)
    assert all("useState" not in h.text for h in hits)  # stale vectors are gone

    service.delete_document(db, doc)
    assert db.query(DocumentChunk).count() == 0
    assert index.search(db, world["paid"].id, query, emb.model_name, emb.dim, 3) == []


# --------------------------------------------------------------------------- FAISS retrieval

def test_faiss_ranks_relevant_chunk_first_and_applies_threshold(db, world):
    _ingest(db, world["paid"], "Hooks", PAID_TEXT + "\n\n" + "Gardening tips: water tomatoes every morning in summer.")
    retrieval = service.retrieve(db, world["paid"].id, "How does the useState hook store a value?")
    assert retrieval.sources and "useState" in retrieval.sources[0].text
    assert all(s.score >= settings.RAG_MIN_SCORE for s in retrieval.sources)
    unrelated = service.retrieve(db, world["paid"].id, "quantum chromodynamics lagrangian")
    assert unrelated.sources == [] and unrelated.context == ""  # nearest neighbours are not evidence


def test_retrieval_never_crosses_courses(db, world):
    _ingest(db, world["paid"], "Paid", PAID_TEXT)
    _ingest(db, world["free"], "Free", FREE_TEXT)
    hits = service.retrieve(db, world["free"].id, "react hooks useState useEffect capstone rubric").sources
    assert all("React" not in s.text and "useState" not in s.text for s in hits)
    paid_hits = service.retrieve(db, world["paid"].id, "python generators yield lazily").sources
    assert all("generators" not in s.text for s in paid_hits)


def test_context_dedupes_and_respects_budget(db, world, monkeypatch):
    text = "\n\n".join(["Repeated paragraph about hooks and state management in React apps."] * 4
                       + [f"Unique paragraph {i} about hooks state effects and rendering in React." for i in range(8)])
    _ingest(db, world["paid"], "Dup", text)
    monkeypatch.setattr(settings, "RAG_CHUNK_CHARS", 120)
    monkeypatch.setattr(settings, "RAG_MAX_CONTEXT_CHARS", 400)
    sources = service.retrieve(db, world["paid"].id, "hooks state effects React", top_k=8).sources
    texts = [s.text.strip().lower() for s in sources]
    assert len(texts) == len(set(texts))
    assert sum(len(s.text) for s in sources) <= 400 + max(len(s.text) for s in sources)


def test_index_rebuilds_after_cache_loss_and_detects_new_documents(db, world):
    _ingest(db, world["paid"], "Hooks", PAID_TEXT)
    emb = embedder.get_embedder()
    q = emb.encode(["useEffect side effects"])[0]
    first = index.search(db, world["paid"].id, q, emb.model_name, emb.dim, 3)
    index.clear()  # simulate a process restart / redeploy: no index in memory, only the database
    again = index.search(db, world["paid"].id, q, emb.model_name, emb.dim, 3)
    assert [h.chunk_id for h in first] == [h.chunk_id for h in again] and first
    _ingest(db, world["paid"], "More", "useEffect cleanup functions prevent memory leaks in effects.")
    after = index.search(db, world["paid"].id, q, emb.model_name, emb.dim, 5)
    assert any("cleanup" in h.text for h in after)  # fingerprint changed -> rebuilt, no stale index


def test_documents_from_another_embedding_model_are_not_searched(db, world):
    doc, _ = _ingest(db, world["paid"], "Hooks", PAID_TEXT)
    doc.embedding_model = "some/other-model"
    db.commit()
    emb = embedder.get_embedder()
    assert index.search(db, world["paid"].id, emb.encode(["useState"])[0], emb.model_name, emb.dim, 3) == []


def test_missing_faiss_is_an_operational_error(db, world, monkeypatch):
    _ingest(db, world["paid"], "Hooks", PAID_TEXT)

    def no_faiss():
        raise index.IndexUnavailable("faiss-cpu is not installed")
    monkeypatch.setattr(index, "_faiss", no_faiss)
    index.clear()
    with pytest.raises(index.IndexUnavailable):
        service.retrieve(db, world["paid"].id, "useState")


def test_missing_embedding_model_is_reported_not_faked(monkeypatch):
    embedder.set_embedder(None)

    class Broken:
        def __init__(self, name):
            raise embedder.EmbeddingUnavailable("could not be loaded")
    monkeypatch.setattr(embedder, "SentenceTransformerEmbedder", Broken)
    assert embedder.health()["status"] == "unavailable"
    with pytest.raises(embedder.EmbeddingUnavailable):
        embedder.embed_texts(["hello"])


def test_real_sentence_transformer_embeds_and_ranks():
    embedder.set_embedder(None)
    try:
        model = embedder.get_embedder()
    except embedder.EmbeddingUnavailable as exc:
        pytest.skip(f"Sentence Transformers model unavailable in this environment: {exc}")
    texts = ["How do I reverse a list in Python?", "Use list.reverse() or slicing [::-1].", "Bananas are yellow."]
    vectors = model.encode(texts)
    assert vectors.shape == (3, model.dim) and abs(float(np.linalg.norm(vectors[0])) - 1) < 1e-3
    assert float(vectors[0] @ vectors[1]) > float(vectors[0] @ vectors[2])


# --------------------------------------------------------------------------- authorization (unit)

def test_authorize_matrix(db, world):
    w = world
    with pytest.raises(RagAccessDenied) as e:
        authorize_course_access(db, w["student"], w["paid"].id)
    assert e.value.code == "NOT_ENROLLED"
    enroll(db, w["student"], w["paid"])  # enrollment alone is not a purchase
    with pytest.raises(RagAccessDenied) as e:
        authorize_course_access(db, w["student"], w["paid"].id)
    assert e.value.code == "PAYMENT_REQUIRED"
    db.add(Payment(user_id=w["student"].id, course_id=w["paid"].id, razorpay_order_id="order_UNITTEST0001",
                   amount=1, currency="INR", status=PaymentStatus.PENDING.value))
    db.commit()
    with pytest.raises(RagAccessDenied):  # pending is not paid
        authorize_course_access(db, w["student"], w["paid"].id)
    db.query(Payment).update({"status": PaymentStatus.SUCCESS.value})
    db.commit()
    assert authorize_course_access(db, w["student"], w["paid"].id).id == w["paid"].id
    assert authorize_course_access(db, w["instructor"], w["paid"].id)
    assert authorize_course_access(db, w["admin"], w["draft"].id)
    with pytest.raises(RagAccessDenied) as e:
        authorize_course_access(db, w["other_instructor"], w["paid"].id)
    assert e.value.status_code == 403
    with pytest.raises(RagAccessDenied) as e:  # hidden course looks identical to a missing one
        authorize_course_access(db, w["student"], w["draft"].id)
    assert e.value.status_code == 404
    with pytest.raises(RagAccessDenied) as e:
        authorize_course_access(db, w["student"], 99999)
    assert e.value.status_code == 404


# --------------------------------------------------------------------------- /api/ai/chat end to end

def chat(client, user, course_id, message="How does the useState hook store a value?", **extra):
    return client.post("/api/ai/chat", json={"message": message, "course_id": course_id, **extra},
                       headers=auth(user) if user else {})


def test_unauthenticated_chat_is_rejected(client, world, recorder):
    assert chat(client, None, world["paid"].id).status_code in (401, 403)
    assert recorder.calls == []


def test_unpaid_user_cannot_extract_paid_course_material(client, db, world, recorder):
    _ingest(db, world["paid"], "Secret", PAID_TEXT)
    for attempt in (
        dict(course_id=world["paid"].id),                                  # knows the course id
        dict(course_id=str(world["paid"].id)),                             # coerced id
        dict(course_id=world["paid"].id, message="Premium secret capstone rubric custom hooks"),
    ):
        res = chat(client, world["stranger"], **{"course_id": attempt["course_id"],
                                                 "message": attempt.get("message", "useState hook")})
        assert res.status_code == 403, res.text
        assert "capstone" not in res.text and "useState" not in res.text
    assert recorder.calls == []  # the LLM never saw anything; retrieval never ran


def test_enrolled_without_verified_payment_is_denied(client, db, world, recorder):
    _ingest(db, world["paid"], "Secret", PAID_TEXT)
    enroll(db, world["student"], world["paid"])  # e.g. a forged/direct enrollment row
    res = chat(client, world["student"], world["paid"].id)
    assert res.status_code == 403 and res.json()["code"] == "PAYMENT_REQUIRED"
    assert recorder.calls == []


def test_paid_student_gets_grounded_answer_with_real_sources(client, db, world, recorder):
    grant_paid_access(db, world["student"], world["paid"])
    res = chat(client, world["student"], world["paid"].id)
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["grounded"] is True and body["sources"]
    assert body["sources"][0]["cited"] is True and body["sources"][0]["title"]
    prompt = recorder.calls[0]["prompt"]
    assert "<course_material>" in prompt and "useState" in prompt
    assert "generators" not in prompt  # nothing from the other course
    assert "COURSE MATERIAL RULES" in recorder.calls[0]["system_instruction"]


def test_refund_revokes_access_immediately_even_after_a_successful_answer(client, db, world, recorder):
    grant_paid_access(db, world["student"], world["paid"])
    assert chat(client, world["student"], world["paid"].id).status_code == 200  # index + docs now cached
    db.query(Payment).update({"status": PaymentStatus.REFUNDED.value})
    db.query(Enrollment).update({"is_active": False})
    db.commit()
    res = chat(client, world["student"], world["paid"].id)
    assert res.status_code == 403
    assert len(recorder.calls) == 1  # no second LLM call, nothing served from a warm cache


def test_free_course_requires_enrollment(client, db, world, recorder):
    assert chat(client, world["student"], world["free"].id, "python generators yield").status_code == 403
    enroll(db, world["student"], world["free"])
    res = chat(client, world["student"], world["free"].id, "python generators yield lazily")
    assert res.status_code == 200 and res.json()["grounded"] is True


def test_enrollment_in_one_course_does_not_unlock_another(client, db, world, recorder):
    enroll(db, world["student"], world["free"])
    res = chat(client, world["student"], world["paid"].id)
    assert res.status_code == 403
    assert recorder.calls == []


def test_unpublished_course_is_hidden_from_students(client, db, world, recorder):
    res = chat(client, world["student"], world["draft"].id)
    assert res.status_code == 404 and "secret syllabus" not in res.text


def test_owner_instructor_allowed_other_instructor_denied(client, db, world, recorder):
    assert chat(client, world["instructor"], world["paid"].id).status_code == 200
    assert chat(client, world["other_instructor"], world["paid"].id).status_code == 403


def test_v3_tutor_route_enforces_the_same_authorization(client, db, world, recorder):
    res = client.post("/api/ai/tutor/chat", params={"message": "useState hook", "session_token": "s1",
                                                    "course_id": world["paid"].id}, headers=auth(world["stranger"]))
    assert res.status_code == 403 and recorder.calls == []
    grant_paid_access(db, world["student"], world["paid"])
    ok = client.post("/api/ai/tutor/chat", params={"message": "useState hook", "session_token": "s1",
                                                   "course_id": world["paid"].id}, headers=auth(world["student"]))
    assert ok.status_code == 200 and ok.json()["grounded"] is True


def test_no_relevant_context_is_admitted_not_disguised(client, db, world, recorder):
    grant_paid_access(db, world["student"], world["paid"])
    recorder.reply = "I could not find this in the course material. General explanation: ..."
    res = chat(client, world["student"], world["paid"].id, "quantum chromodynamics lagrangian")
    body = res.json()
    assert res.status_code == 200 and body["grounded"] is False and body["sources"] == []
    assert "no relevant course material was found" in recorder.calls[0]["system_instruction"]
    assert "<course_material>" not in recorder.calls[0]["prompt"]


def test_invented_citations_are_removed(client, db, world, recorder):
    grant_paid_access(db, world["student"], world["paid"])
    recorder.reply = "Real point [S1]. Made-up point [S9] and [S42]."
    body = chat(client, world["student"], world["paid"].id).json()
    assert "[S9]" not in body["response"] and "[S42]" not in body["response"] and "[S1]" in body["response"]
    assert {s["id"] for s in body["sources"]} <= {1, 2, 3, 4, 5}


def test_greeting_skips_retrieval_but_still_requires_authorization(client, db, world, recorder, monkeypatch):
    calls = []
    monkeypatch.setattr(service, "retrieve", lambda *a, **k: calls.append(1))
    assert chat(client, world["stranger"], world["paid"].id, "hello").status_code == 403
    grant_paid_access(db, world["student"], world["paid"])
    res = chat(client, world["student"], world["paid"].id, "thanks")
    assert res.status_code == 200 and calls == [] and res.json()["grounded"] is False


def test_embedding_outage_returns_503_without_leaking_or_calling_llm(client, db, world, recorder, monkeypatch):
    grant_paid_access(db, world["student"], world["paid"])
    embedder.set_embedder(None)

    class Broken:
        def __init__(self, name):
            raise embedder.EmbeddingUnavailable("secret internal path /models/x")
    monkeypatch.setattr(embedder, "SentenceTransformerEmbedder", Broken)
    res = chat(client, world["student"], world["paid"].id)
    assert res.status_code == 503 and "secret internal path" not in res.text
    assert recorder.calls == []


def test_chat_without_course_still_works_as_a_general_tutor(client, world, recorder):
    res = client.post("/api/ai/chat", json={"message": "Explain recursion"}, headers=auth(world["stranger"]))
    assert res.status_code == 200 and res.json()["grounded"] is False
    assert "COURSE MATERIAL RULES" not in recorder.calls[0]["system_instruction"]


@pytest.mark.parametrize("kind,status", [(llm.TIMEOUT, 504), (llm.RATE_LIMITED, 429), (llm.EMPTY, 502),
                                         (llm.NOT_CONFIGURED, 503), (llm.BLOCKED, 422)])
def test_llm_failures_become_http_errors_and_are_not_saved(client, db, world, monkeypatch, kind, status):
    from models.analytics import ChatHistory

    def fail(prompt, **kw):
        raise llm.LLMError(kind)
    monkeypatch.setattr(llm, "generate", fail)
    res = client.post("/api/ai/chat", json={"message": "Explain recursion"}, headers=auth(world["stranger"]))
    assert res.status_code == status and res.json()["detail"]
    assert db.query(ChatHistory).count() == 0


# --------------------------------------------------------------------------- upload endpoints

@pytest.fixture
def bg_db(monkeypatch):
    import routes.rag as rag_routes
    monkeypatch.setattr(rag_routes, "SessionLocal", TestingSessionLocal)


def upload(client, user, course_id, name="notes.txt", data=PAID_TEXT.encode()):
    return client.post(f"/api/rag/courses/{course_id}/documents", files={"file": (name, data, "text/plain")},
                       headers=auth(user))


def test_owner_can_upload_and_status_is_tracked(client, db, world, bg_db):
    res = upload(client, world["instructor"], world["paid"].id)
    assert res.status_code == 202 and res.json()["duplicate"] is False
    docs = client.get(f"/api/rag/courses/{world['paid'].id}/documents", headers=auth(world["instructor"])).json()
    assert docs[0]["status"] == "ready" and docs[0]["chunk_count"] > 0 and docs[0]["error"] is None
    again = upload(client, world["instructor"], world["paid"].id)
    assert again.json()["duplicate"] is True
    assert len(client.get(f"/api/rag/courses/{world['paid'].id}/documents", headers=auth(world["instructor"])).json()) == 1


def test_bad_uploads_are_rejected_with_reasons(client, db, world, bg_db):
    assert upload(client, world["instructor"], world["paid"].id, "x.exe", b"MZ").status_code == 422
    assert upload(client, world["instructor"], world["paid"].id, "x.txt", b"").status_code == 422
    assert upload(client, world["instructor"], world["paid"].id, "x.pdf", b"not pdf").status_code == 422


def test_students_and_other_instructors_cannot_manage_documents(client, db, world, bg_db):
    assert upload(client, world["student"], world["paid"].id).status_code == 403
    assert upload(client, world["other_instructor"], world["paid"].id).status_code == 404
    upload(client, world["instructor"], world["paid"].id)
    doc = db.query(CourseDocument).first()
    for user in (world["student"], world["other_instructor"]):
        assert client.delete(f"/api/rag/documents/{doc.id}", headers=auth(user)).status_code in (403, 404)
        assert client.post(f"/api/rag/documents/{doc.id}/reindex", headers=auth(user)).status_code in (403, 404)
    assert client.get(f"/api/rag/courses/{world['paid'].id}/documents", headers=auth(world["student"])).status_code == 403
    assert db.query(CourseDocument).count() == 1


def test_reindex_and_delete_by_owner(client, db, world, bg_db):
    upload(client, world["instructor"], world["paid"].id)
    doc = db.query(CourseDocument).first()
    assert client.post(f"/api/rag/documents/{doc.id}/reindex", headers=auth(world["instructor"])).status_code == 202
    assert client.delete(f"/api/rag/documents/{doc.id}", headers=auth(world["instructor"])).status_code == 204
    assert db.query(CourseDocument).count() == 0 and db.query(DocumentChunk).count() == 0


def test_uploaded_document_is_retrievable_only_after_purchase(client, db, world, bg_db, recorder):
    upload(client, world["instructor"], world["paid"].id, data=b"The final exam covers monads and functors in depth.")
    assert chat(client, world["stranger"], world["paid"].id, "what does the final exam cover monads").status_code == 403
    grant_paid_access(db, world["student"], world["paid"])
    res = chat(client, world["student"], world["paid"].id, "what does the final exam cover monads functors")
    assert res.status_code == 200 and "monads" in recorder.calls[-1]["prompt"]


def test_generic_overview_question_gets_the_course_outline(client, db, world, recorder):
    grant_paid_access(db, world["student"], world["paid"])
    res = chat(client, world["student"], world["paid"].id, "What will I learn in this course?")
    body = res.json()
    assert res.status_code == 200 and body["grounded"] is True and body["sources"]
    assert "React Masterclass" in recorder.calls[0]["prompt"]
    # ...but a specific off-topic question still gets no material (the fallback is only for overviews)
    res = chat(client, world["student"], world["paid"].id, "quantum chromodynamics lagrangian")
    assert res.json()["grounded"] is False


def test_overview_detection():
    assert service.is_overview_query("what is covered in this course?")
    assert service.is_overview_query("Give me the syllabus")
    assert not service.is_overview_query("How do hooks work?")


def test_grouped_citations_are_validated_per_source():
    sources = [service.Source(n=1, document_id=1, title="A", section=None, score=0.5, text="x"),
               service.Source(n=2, document_id=1, title="A", section=None, score=0.4, text="y")]
    text, public = service.resolve_citations("Fact [S1, S2]. Bogus [S7]. Mixed [S2, S9]. Single [S1].", sources)
    assert text == "Fact [S1, S2]. Bogus . Mixed [S2]. Single [S1]."
    assert [s["cited"] for s in public] == [True, True]
    _, public = service.resolve_citations("No citations here.", sources)
    assert [s["cited"] for s in public] == [False, False]
